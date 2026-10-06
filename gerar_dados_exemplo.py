#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera planilhas .xlsx 100% FICTICIAS para testar o dashboard.

Colunas (mesmas esperadas pelo painel):
    Data | Departamento | Descriçao das Naturezas | Favorecido | Beneficiario | Banco | Valor

Uso:
    python gerar_dados_exemplo.py            # grava em dados_exemplo/
    python gerar_dados_exemplo.py --saida bd_exemplo

Requisito: pip install openpyxl
"""
import argparse
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

SEED = 2026
CABECALHO = ["Data", "Departamento", "Descriçao das Naturezas", "Favorecido",
             "Beneficiario", "Banco", "Valor"]
DEPARTAMENTOS = ["Financeiro", "Operações", "TI", "Marketing", "Recursos Humanos",
                 "Comercial", "Jurídico", "Logística"]
NATUREZAS = ["Serviços de terceiros", "Material de escritório", "Licenças de software",
             "Manutenção predial", "Energia elétrica", "Telefonia e internet",
             "Publicidade", "Fretes e transportes", "Consultoria", "Aluguéis",
             "Treinamentos", "Viagens e hospedagem"]
BANCOS = ["Banco A", "Banco B", "Banco C", "Banco D"]
EMPRESAS = [f"Fornecedor {i:03d}" for i in range(1, 61)] + \
           ["Empresa Alfa", "Empresa Beta", "Empresa Gama", "Empresa Delta", "Empresa Ômega"]
# alguns favorecidos agrupam varios beneficiarios (ex.: cobranca centralizada)
FAVORECIDOS_GRUPO = ["Cobrança Central Exemplo", "Cooperativa Exemplo", "Factoring Exemplo"]

INICIO = date(2026, 8, 3)
FIM = date(2026, 10, 2)


def dias_uteis(ini, fim):
    d = ini
    while d <= fim:
        if d.weekday() < 5:
            yield d
        d += timedelta(days=1)


def gera(n_linhas, rnd):
    dias = list(dias_uteis(INICIO, FIM))
    # fornecedores recorrentes pagam com mais frequencia
    pesos = [rnd.choice([1, 1, 2, 4, 8]) for _ in EMPRESAS]
    linhas = []
    for _ in range(n_linhas):
        d = rnd.choice(dias)
        if rnd.random() < 0.02:                      # raros pagamentos em fim de semana
            d += timedelta(days=5 - d.weekday() + rnd.choice([0, 1]))
        emp = rnd.choices(EMPRESAS, weights=pesos)[0]
        valor = round(min(max(rnd.lognormvariate(7.2, 1.15), 35.0), 48000.0), 2)
        fav = emp if rnd.random() > 0.12 else rnd.choice(FAVORECIDOS_GRUPO)
        linhas.append([d, rnd.choice(DEPARTAMENTOS), rnd.choice(NATUREZAS),
                       fav, emp, rnd.choice(BANCOS), valor])
    linhas.sort(key=lambda r: r[0])
    return linhas


def grava(caminho, linhas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    ws.append(CABECALHO)
    for r in linhas:
        ws.append(r)
    for linha in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        linha[0].number_format = "DD/MM/YYYY"
    for linha in ws.iter_rows(min_row=2, min_col=7, max_col=7):
        linha[0].number_format = "#,##0.00"
    for col, larg in zip("ABCDEFG", (12, 18, 26, 28, 24, 12, 14)):
        ws.column_dimensions[col].width = larg
    wb.save(caminho)


def main():
    ap = argparse.ArgumentParser(description="Gera planilhas de exemplo (dados ficticios)")
    ap.add_argument("--saida", default="dados_exemplo")
    ap.add_argument("--linhas", type=int, default=1500)
    args = ap.parse_args()
    rnd = random.Random(SEED)
    pasta = Path(args.saida)
    pasta.mkdir(parents=True, exist_ok=True)
    linhas = gera(args.linhas, rnd)
    destino = pasta / "Pagamentos_Exemplo_2026.xlsx"
    grava(destino, linhas)
    total = sum(r[6] for r in linhas)
    print(f"{destino}: {len(linhas)} linhas, total fictício {total:,.2f}")


if __name__ == "__main__":
    main()